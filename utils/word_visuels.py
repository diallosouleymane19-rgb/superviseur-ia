# -*- coding: utf-8 -*-
"""
Visuels des exports Word : bloc d'indicateurs clés et graphiques (images PNG au format français).
Les graphiques Plotly de l'écran ne peuvent pas être convertis en image sur le serveur :
ils sont redessinés ici avec matplotlib, avec les mêmes données.
SMD Global Consulting LLC
"""
import io
import textwrap

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from utils.sig_pcg import nb_fr

# Couleurs : une seule teinte bleue pour les montants (clair = N-1 ou secondaire, foncé = N ou principal)
COULEUR_N1, COULEUR_N = "#86b6ef", "#1c5cab"
COULEUR_ORANGE = "#eb6834"          # seconde série d'une autre nature (ex. emplois face aux ressources)
COULEUR_REFERENCE = "#52514e"       # valeur attendue / de référence (trait)
GRIS = RGBColor(0x52, 0x51, 0x4E)
VERT, ROUGE = RGBColor(0x0B, 0x7A, 0x0B), RGBColor(0xC0, 0x29, 0x29)
TONS = {"bon": VERT, "mauvais": ROUGE, "neutre": GRIS, None: GRIS}


# ─── Outils Word ─────────────────────────────────────────────────────────────

def fond(cellule, couleur_hex):
    tc_pr = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur_hex)
    tc_pr.append(shd)


def texte(cellule, contenu, gras=False, taille=9.5, couleur=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    from utils.export_word import typo_fr
    p = cellule.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(typo_fr(str(contenu)))
    r.bold = gras or None
    r.font.size = Pt(taille)
    if couleur is not None:
        r.font.color.rgb = couleur
    return r


def largeurs(t, valeurs):
    """Largeurs de colonnes fixes (Word et LibreOffice)."""
    t.autofit = False
    for j, w in enumerate(valeurs):
        t.columns[j].width = w
        for cel in t.columns[j].cells:
            cel.width = w


def garder_ensemble(t):
    """Le tableau ne se coupe pas entre deux pages s'il tient sur une seule."""
    for i, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if i < len(t.rows) - 1:
            for cel in row.cells:
                for par in cel.paragraphs:
                    par.paragraph_format.keep_with_next = True


# ─── Indicateurs clés ────────────────────────────────────────────────────────

def bloc_indicateurs(doc, indicateurs, par_ligne=None):
    """indicateurs : liste de dict {libelle, valeur (texte), detail (texte, facultatif),
    ton ('bon' | 'mauvais' | 'neutre', facultatif)}. Lignes équilibrées : 5 cases au plus sur une ligne,
    sinon 4 au plus par ligne réparties à parts égales (6 → 3 + 3, 7 → 4 + 3, 9 → 3 + 3 + 3)."""
    import math
    n = len(indicateurs)
    if par_ligne is None:
        par_ligne = n if n <= 5 else math.ceil(n / math.ceil(n / 4))
    for k in range(0, n, par_ligne):
        groupe = indicateurs[k:k + par_ligne]
        avec_detail = any(i.get("detail") for i in groupe)
        t = doc.add_table(rows=3 if avec_detail else 2, cols=len(groupe))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, ind in enumerate(groupe):
            for i in range(len(t.rows)):
                fond(t.cell(i, j), "F2F5FA")
            texte(t.cell(0, j), ind["libelle"], taille=8.5, couleur=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER)
            texte(t.cell(1, j), ind["valeur"], gras=True, taille=13, align=WD_ALIGN_PARAGRAPH.CENTER)
            if avec_detail:
                texte(t.cell(2, j), ind.get("detail") or "", taille=8, couleur=TONS.get(ind.get("ton"), GRIS),
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        garder_ensemble(t)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)


def inserer_graphique(doc, png, largeur_cm=16):
    doc.add_picture(io.BytesIO(png), width=Cm(largeur_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


# ─── Graphiques ──────────────────────────────────────────────────────────────

def _axes(titre, hauteur=3.4):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})
    fig, ax = plt.subplots(figsize=(7.2, hauteur), dpi=200)
    ax.grid(axis="y", color="#e5e5e5", linewidth=0.6)
    ax.set_axisbelow(True)
    for cote in ("top", "right", "left"):
        ax.spines[cote].set_visible(False)
    ax.tick_params(axis="y", length=0, labelsize=7.5, colors="#52514e")
    ax.set_title(titre, fontsize=10, loc="left", color="#0b0b0b")
    return fig, ax


def _png(fig):
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def _etiquettes(ax, barres, fmt=nb_fr, taille=6.5):
    for r in barres:
        h = r.get_height()
        ax.annotate(fmt(h), (r.get_x() + r.get_width() / 2, h), xytext=(0, 3 if h >= 0 else -10),
                    textcoords="offset points", ha="center", fontsize=taille, color="#333333")


def _axe_x(ax, libelles, largeur=14):
    ax.set_xticks(list(range(len(libelles))))
    ax.set_xticklabels([textwrap.fill(str(x), largeur, break_long_words=False) for x in libelles], fontsize=7.5)


def _axe_euros(ax, unite="Montant (€)"):
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: nb_fr(v)))
    ax.set_ylabel(unite, fontsize=8, color="#52514e")
    ax.axhline(0, color="#999999", linewidth=0.8)


def barres_groupees(libelles, vals_a, vals_b, nom_a, nom_b, titre, couleurs=(COULEUR_N1, COULEUR_N)):
    """Deux séries côte à côte (ex. N-1 / N). Retourne un PNG."""
    fig, ax = _axes(titre, 3.6)
    l = 0.38
    x = range(len(libelles))
    b1 = ax.bar([i - l / 2 - 0.01 for i in x], vals_a, l, color=couleurs[0], label=nom_a)
    b2 = ax.bar([i + l / 2 + 0.01 for i in x], vals_b, l, color=couleurs[1], label=nom_b)
    _etiquettes(ax, b1)
    _etiquettes(ax, b2)
    _axe_x(ax, libelles)
    _axe_euros(ax)
    ax.legend(frameon=False, fontsize=8, ncol=2, loc="upper right", bbox_to_anchor=(1, 1.13))
    return _png(fig)


def barres_simples(libelles, valeurs, titre, unite="Montant (€)", fmt=nb_fr):
    """Une série ; les valeurs négatives en orange (repérées aussi par le signe de l'étiquette)."""
    fig, ax = _axes(titre)
    couleurs = [COULEUR_N if v >= 0 else COULEUR_ORANGE for v in valeurs]
    b = ax.bar(range(len(libelles)), valeurs, 0.6, color=couleurs)
    _etiquettes(ax, b, fmt=fmt)
    _axe_x(ax, libelles)
    _axe_euros(ax, unite)
    return _png(fig)


def barres_empilees(colonnes, segments, titre):
    """Barres empilées (parties d'un tout). colonnes : libellés des barres ;
    segments : liste de (nom, [valeur par colonne], couleur). Valeurs ≥ 0."""
    fig, ax = _axes(titre, 3.8)
    bas = [0.0] * len(colonnes)
    total = max(sum(v[i] for _, v, _ in segments) for i in range(len(colonnes))) or 1
    for nom, vals, coul in segments:
        b = ax.bar(range(len(colonnes)), vals, 0.5, bottom=bas, color=coul,
                   label=nom, edgecolor="white", linewidth=1.2)
        fonce = coul.lower() in ("#1c5cab", "#4f8fdc", "#c2410c", "#eb6834")
        for r, v, y0 in zip(b, vals, bas):
            if v > total * 0.04:   # étiquette seulement si le segment est assez haut pour la contenir
                ax.text(r.get_x() + r.get_width() / 2, y0 + v / 2, nb_fr(v), ha="center", va="center",
                        fontsize=6.5, color="white" if fonce else "#0b0b0b")
        bas = [x + v for x, v in zip(bas, vals)]
    _axe_x(ax, colonnes, 20)
    _axe_euros(ax)
    ax.legend(frameon=False, fontsize=7, loc="center left", bbox_to_anchor=(1.0, 0.5))
    return _png(fig)


def barres_et_courbe(libelles, barres, courbe, titre):
    """barres : liste de (nom, valeurs, couleur) côte à côte ; courbe : (nom, valeurs) en trait."""
    fig, ax = _axes(titre, 3.6)
    n = len(barres)
    l = 0.8 / max(n, 1)
    for k, (nom, vals, coul) in enumerate(barres):
        b = ax.bar([i + (k - (n - 1) / 2) * l for i in range(len(libelles))], vals, l * 0.95, color=coul, label=nom)
        _etiquettes(ax, b)
    if courbe:
        nom, vals = courbe
        ax.plot(range(len(libelles)), vals, color=COULEUR_REFERENCE, linewidth=2, marker="o", markersize=5,
                label=nom)
    _axe_x(ax, libelles)
    _axe_euros(ax)
    ax.legend(frameon=False, fontsize=8, ncol=3, loc="upper right", bbox_to_anchor=(1, 1.14))
    return _png(fig)


def benford(chiffres, observe_pct, attendu_pct, titre="Premier chiffre : fréquences observées et attendues"):
    """Barres = fréquences observées (%), trait = loi de Benford (%)."""
    fig, ax = _axes(titre)
    b = ax.bar(range(len(chiffres)), observe_pct, 0.6, color=COULEUR_N, label="Observé")
    ax.plot(range(len(chiffres)), attendu_pct, color=COULEUR_ORANGE, linewidth=2, marker="o", markersize=5,
            label="Attendu (loi de Benford)")
    _etiquettes(ax, b, fmt=lambda v: nb_fr(v, 1))
    _axe_x(ax, [str(c) for c in chiffres])
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: nb_fr(v, 0) + " %"))
    ax.set_xlabel("Premier chiffre", fontsize=8, color="#52514e")
    ax.legend(frameon=False, fontsize=8, ncol=2, loc="upper right")
    return _png(fig)
