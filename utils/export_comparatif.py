# -*- coding: utf-8 -*-
"""
Export Word du Comparatif N / N-1 : indicateurs clés, graphiques, tableaux et analyse des écarts.
SMD Global Consulting LLC - PCG France
"""
import io
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from utils.export_word import document_smd, typo_fr
from utils.sig_pcg import nb_fr, nb_fr_signe

# N-1 clair, N foncé : même teinte (même grandeur à deux dates)
COULEUR_N1, COULEUR_N = "#86b6ef", "#1c5cab"
BLEU_TITRE = RGBColor(0x1F, 0x4E, 0x79)
GRIS = RGBColor(0x52, 0x51, 0x4E)
VERT, ROUGE = RGBColor(0x0B, 0x7A, 0x0B), RGBColor(0xC0, 0x29, 0x29)


# ─── Petits outils Word ──────────────────────────────────────────────────────

def _fond(cellule, couleur_hex):
    tc_pr = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur_hex)
    tc_pr.append(shd)


def _texte(cellule, texte, gras=False, taille=9.5, couleur=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cellule.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(typo_fr(str(texte)))
    r.bold = gras or None
    r.font.size = Pt(taille)
    if couleur is not None:
        r.font.color.rgb = couleur
    return r


def _titre(doc, texte, niveau=1):
    h = doc.add_heading(level=niveau)
    h.add_run(typo_fr(texte))
    h.paragraph_format.keep_with_next = True
    return h


def _note(doc, texte):
    p = doc.add_paragraph()
    r = p.add_run(typo_fr(texte))
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GRIS
    return p


def _ecart_txt(ea, ep):
    return f"{nb_fr_signe(ea)} € ({nb_fr_signe(ep, 1)} %)"


# ─── Indicateurs clés ────────────────────────────────────────────────────────

def _bloc_kpi(doc, kpis):
    """kpis : liste de (libellé, valeur N, écart €, écart %, hausse_favorable)."""
    t = doc.add_table(rows=3, cols=len(kpis))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (lib, vn, ea, ep, hausse_ok) in enumerate(kpis):
        for i in range(3):
            _fond(t.cell(i, j), "F2F5FA")
        _texte(t.cell(0, j), lib, taille=9, couleur=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER)
        _texte(t.cell(1, j), f"{nb_fr(vn)} €", gras=True, taille=14, align=WD_ALIGN_PARAGRAPH.CENTER)
        if ea == 0:
            couleur, fleche = GRIS, "="
        else:
            favorable = (ea > 0) == hausse_ok
            couleur, fleche = (VERT if favorable else ROUGE), ("▲" if ea > 0 else "▼")
        _texte(t.cell(2, j), f"{fleche} {_ecart_txt(ea, ep)}", taille=8.5, couleur=couleur,
               align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()


# ─── Tableaux ────────────────────────────────────────────────────────────────

def _tableau(doc, df, label_n1, label_n):
    """Tableau comparatif : Rubrique · N-1 · N · Écart (€) · Écart (%) ; totaux et résultats en gras."""
    cols = ["Rubrique", label_n1, label_n, "Écart (€)", "Écart (%)"]
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, c in enumerate(cols):
        _fond(t.rows[0].cells[j], "1F4E79")
        _texte(t.rows[0].cells[j], c, gras=True, taille=9, couleur=RGBColor(0xFF, 0xFF, 0xFF),
               align=WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT)
    for _, ligne in df.iterrows():
        rub = str(ligne["Rubrique"])
        total = rub.upper().startswith("TOTAL") or rub.startswith("Résultat")
        cells = t.add_row().cells
        valeurs = [rub, nb_fr(ligne[label_n1]), nb_fr(ligne[label_n]),
                   nb_fr_signe(ligne["Écart (€)"]), nb_fr_signe(ligne["Écart (%)"], 1)]
        for j, v in enumerate(valeurs):
            if total:
                _fond(cells[j], "EEF2F7")
            _texte(cells[j], v, gras=total, taille=9,
                   align=WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT)
    _largeurs(t, [Cm(6.4), Cm(2.5), Cm(2.5), Cm(2.5), Cm(2.1)])
    _garder_ensemble(t)
    doc.add_paragraph()


def _largeurs(t, largeurs):
    """Largeurs de colonnes fixes (Word et LibreOffice)."""
    t.autofit = False
    for j, w in enumerate(largeurs):
        t.columns[j].width = w
        for cel in t.columns[j].cells:
            cel.width = w


def _garder_ensemble(t):
    """Le tableau ne se coupe pas entre deux pages s'il tient sur une seule."""
    for i, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
        if i < len(t.rows) - 1:
            for cel in row.cells:
                for par in cel.paragraphs:
                    par.paragraph_format.keep_with_next = True


# ─── Graphiques (images PNG) ─────────────────────────────────────────────────

def graphique_barres(libelles, vals_n1, vals_n, label_n1, label_n, titre):
    """Barres groupées N-1 / N, valeurs au format français. Retourne un PNG (bytes)."""
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})
    fig, ax = plt.subplots(figsize=(7.2, 3.6), dpi=200)
    x = range(len(libelles))
    l = 0.38
    b1 = ax.bar([i - l / 2 - 0.01 for i in x], vals_n1, l, color=COULEUR_N1, label=label_n1)
    b2 = ax.bar([i + l / 2 + 0.01 for i in x], vals_n, l, color=COULEUR_N, label=label_n)
    for barres in (b1, b2):
        for r in barres:
            h = r.get_height()
            ax.annotate(nb_fr(h), (r.get_x() + r.get_width() / 2, h), xytext=(0, 3 if h >= 0 else -10),
                        textcoords="offset points", ha="center", fontsize=6.5, color="#333333")
    ax.set_xticks(list(x))
    import textwrap
    ax.set_xticklabels([textwrap.fill(str(x_), 14) for x_ in libelles], fontsize=7.5)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: nb_fr(v)))
    ax.set_ylabel("Montant (€)", fontsize=8, color="#52514e")
    ax.axhline(0, color="#999999", linewidth=0.8)
    ax.grid(axis="y", color="#e5e5e5", linewidth=0.6)
    ax.set_axisbelow(True)
    for cote in ("top", "right", "left"):
        ax.spines[cote].set_visible(False)
    ax.tick_params(axis="y", length=0, labelsize=7.5, colors="#52514e")
    ax.set_title(titre, fontsize=10, loc="left", color="#0b0b0b")
    ax.legend(frameon=False, fontsize=8, ncol=2, loc="upper right", bbox_to_anchor=(1, 1.13))
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


# ─── Document complet ────────────────────────────────────────────────────────

def export_word_comparatif(entreprise, label_n, label_n1, kpis_cdr, kpis_bilan, df_sig, df_prod, df_chg,
                           df_actif, df_passif, graph_sig, graph_bilan, alertes):
    """Retourne le document Word (BytesIO). graph_* : (libellés, valeurs N-1, valeurs N)."""
    doc = document_smd()
    t = doc.add_heading(level=0)
    t.add_run(typo_fr(f"Comparatif {label_n} / {label_n1} – {entreprise}"))
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(typo_fr(f"Exercice {label_n} comparé à l'exercice {label_n1} · "
                          f"édité le {datetime.now().strftime('%d/%m/%Y')}"))
    r.font.size = Pt(9)
    r.font.color.rgb = GRIS

    # Synthèse : indicateurs clés et analyse des écarts
    _titre(doc, "Synthèse", 1)
    _bloc_kpi(doc, kpis_cdr)
    _bloc_kpi(doc, kpis_bilan)
    _note(doc, "▲ / ▼ : sens de la variation par rapport à l'exercice précédent ; "
               "vert = évolution favorable, rouge = défavorable (pour le BFR, une hausse est défavorable).")
    _titre(doc, "Analyse des écarts", 2)
    for _, msg in (alertes or [("info", "Aucune alerte significative détectée.")]):
        doc.add_paragraph(typo_fr(msg), style="List Bullet")

    # Compte de résultat
    _titre(doc, "Compte de résultat", 1)
    doc.add_picture(io.BytesIO(graphique_barres(*graph_sig, label_n1, label_n,
                                                "Soldes intermédiaires de gestion")), width=Cm(16))
    _titre(doc, "Soldes intermédiaires de gestion", 2)
    _tableau(doc, df_sig, label_n1, label_n)
    _titre(doc, "Produits", 2)
    _tableau(doc, df_prod, label_n1, label_n)
    _titre(doc, "Charges", 2)
    _tableau(doc, df_chg, label_n1, label_n)

    # Bilan
    _titre(doc, "Bilan", 1)
    doc.add_picture(io.BytesIO(graphique_barres(*graph_bilan, label_n1, label_n,
                                                "Structure du bilan")), width=Cm(16))
    _titre(doc, "Actif", 2)
    _tableau(doc, df_actif, label_n1, label_n)
    _titre(doc, "Passif", 2)
    _tableau(doc, df_passif, label_n1, label_n)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
