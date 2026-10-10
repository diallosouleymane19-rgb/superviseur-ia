# -*- coding: utf-8 -*-
"""
Module TFT PCG France — Méthode indirecte — SMD Global Consulting LLC
Tableau de Flux de Trésorerie conforme modèle OEC
Horizon 1 à 3 exercices comparatifs
"""
from utils.sig_pcg import nb_fr, nb_fr_signe
from utils.formats import est_tableur, TYPES_BALANCE, TYPES_TABLEUR_CSV
from utils.security import sanitize_filename
import streamlit as st
import pandas as pd
import plotly.graph_objects as go
from io import BytesIO
from utils.sig_pcg import eur_fr

# ─── Structure TFT méthode indirecte (ANC / CRC 99-02) ───────────────────────

TFT_STRUCTURE = {
    "I. FLUX DE TRÉSORERIE LIÉS À L'ACTIVITÉ": {
        "color": "#1E8449",
        "lignes": [
            ("Résultat net (bénéfice + / perte -)", "+"),
            ("Dotations aux amortissements et provisions (nettes de reprises)", "+"),
            ("Plus-values de cessions nettes d'impôts", "-"),
            ("Moins-values de cessions nettes d'impôts", "+"),
            ("Variation des stocks (augmentation -)", "±"),
            ("Variation des créances d'exploitation (augmentation -)", "±"),
            ("Variation des dettes d'exploitation (augmentation +)", "±"),
            ("Variation des autres créances (augmentation -)", "±"),
            ("Variation des autres dettes (augmentation +)", "±"),
            ("Impôts sur les sociétés payés", "-"),
            ("Dividendes reçus des participations", "+"),
        ]
    },
    "II. FLUX DE TRÉSORERIE LIÉS AUX OPÉRATIONS D'INVESTISSEMENT": {
        "color": "#2980B9",
        "lignes": [
            ("Acquisitions d'immobilisations incorporelles", "-"),
            ("Acquisitions d'immobilisations corporelles", "-"),
            ("Acquisitions d'immobilisations financières", "-"),
            ("Cessions d'immobilisations incorporelles", "+"),
            ("Cessions d'immobilisations corporelles", "+"),
            ("Cessions d'immobilisations financières", "+"),
            ("Variation des créances sur cessions d'actifs", "±"),
        ]
    },
    "III. FLUX DE TRÉSORERIE LIÉS AUX OPÉRATIONS DE FINANCEMENT": {
        "color": "#8E44AD",
        "lignes": [
            ("Augmentation de capital en numéraire", "+"),
            ("Remboursements de capital", "-"),
            ("Émission d'emprunts", "+"),
            ("Remboursements d'emprunts", "-"),
            ("Dividendes versés aux actionnaires", "-"),
            ("Variation des concours bancaires courants", "±"),
        ]
    }
}

SECTION_KEYS = list(TFT_STRUCTURE.keys())
K_OP  = SECTION_KEYS[0]
K_INV = SECTION_KEYS[1]
K_FIN = SECTION_KEYS[2]


# ─── Calcul ──────────────────────────────────────────────────────────────────

def _calculer_tft(data: dict, exercices: list) -> dict:
    resultats = {}
    for ex in exercices:
        flux_op  = sum(data.get(lib, {}).get(ex, 0) for lib, _ in TFT_STRUCTURE[K_OP]["lignes"])
        flux_inv = sum(data.get(lib, {}).get(ex, 0) for lib, _ in TFT_STRUCTURE[K_INV]["lignes"])
        flux_fin = sum(data.get(lib, {}).get(ex, 0) for lib, _ in TFT_STRUCTURE[K_FIN]["lignes"])
        var_nette = flux_op + flux_inv + flux_fin
        treso_ouv = data.get("Trésorerie à l'ouverture", {}).get(ex, 0)
        treso_clo = treso_ouv + var_nette
        resultats[ex] = {
            "Flux activité (I)": flux_op,
            "Flux investissement (II)": flux_inv,
            "Flux financement (III)": flux_fin,
            "Variation nette (I+II+III)": var_nette,
            "Trésorerie ouverture": treso_ouv,
            "Trésorerie clôture": treso_clo,
        }
    return resultats


def _extraire_tft(df_n, df_n1) -> dict:
    """
    TFT méthode indirecte à partir de deux balances de clôture (N et N-1), avant affectation.
    Retourne {libellé: montant signé} + clés techniques '_treso_ouverture', '_treso_cloture', '_alertes'.
    """
    from utils.sig_pcg import _preparer, _solde
    dn, d1 = _preparer(df_n), _preparer(df_n1)
    c  = lambda d, p, ex=(): _solde(d, p, ex, "credit")
    db = lambda d, p, ex=(): _solde(d, p, ex, "debit")
    var_d = lambda p, ex=(): db(dn, p, ex) - db(d1, p, ex)   # variation d'un solde débiteur
    var_c = lambda p, ex=(): c(dn, p, ex) - c(d1, p, ex)     # variation d'un solde créditeur

    resultat_net = c(dn, ["7"]) - db(dn, ["6"])
    dotations = db(dn, ["681", "686", "687"]) - c(dn, ["781", "786", "787"])
    plus_value = c(dn, ["775"]) - db(dn, ["675"])            # résultat de cession (avant impôt)
    prix_cession = c(dn, ["775"])

    tresorerie = lambda d: db(d, ["50", "51", "53", "54"], ["519"])   # hors concours bancaires
    resultat_n1 = c(d1, ["7"]) - db(d1, ["6"])
    reserves_var = var_c(["106", "11", "12"])
    dividendes = resultat_n1 - reserves_var                   # résultat N-1 non mis en réserve

    capital_var = var_c(["101", "104", "108", "109"])
    emprunts_var = var_c(["16", "17", "455"])               # 455 : comptes courants d'associés
    # Brut sorti lors des cessions = VNC (675) + amortissements cédés
    # amortissements cédés = dotations aux amortissements (6811, 6812) - variation des comptes 28
    amort_cedes = max(db(dn, ["6811", "6812"]) - var_c(["28"]), 0.0)
    brut_cede = db(dn, ["675"]) + amort_cedes
    brut_corp = var_d(["21", "22", "23"]) + brut_cede - var_c(["404", "405"])   # net des dettes sur immobilisations
    brut_incorp = var_d(["20"])
    brut_fin = var_d(["26", "27"])

    v = {
        "Résultat net (bénéfice + / perte -)": resultat_net,
        "Dotations aux amortissements et provisions (nettes de reprises)": dotations,
        "Plus-values de cessions nettes d'impôts": -plus_value if plus_value > 0 else 0.0,
        "Moins-values de cessions nettes d'impôts": -plus_value if plus_value < 0 else 0.0,
        "Variation des stocks (augmentation -)": -var_d(["3"], ["39"]),          # brut : dépréciations dans les dotations
        "Variation des créances d'exploitation (augmentation -)": -var_d(["409", "411", "413", "416", "417", "418"]),  # brut
        "Variation des dettes d'exploitation (augmentation +)": var_c(["401", "403", "408", "419", "42", "43", "445", "447"]),
        "Variation des autres créances (augmentation -)": -var_d(["46", "47", "48"]),
        "Variation des autres dettes (augmentation +)": var_c(["44"], ["445", "447"]),  # dont IS (444)
        "Impôts sur les sociétés payés": 0.0,
        "Acquisitions d'immobilisations incorporelles": -max(brut_incorp, 0.0),
        "Acquisitions d'immobilisations corporelles": -max(brut_corp, 0.0),
        "Acquisitions d'immobilisations financières": -max(brut_fin, 0.0),
        "Cessions d'immobilisations corporelles": prix_cession,
        "Augmentation de capital en numéraire": max(capital_var, 0.0),
        "Remboursements de capital": min(capital_var, 0.0),
        "Émission d'emprunts": max(emprunts_var, 0.0),
        "Remboursements d'emprunts": min(emprunts_var, 0.0),
        "Dividendes versés aux actionnaires": -dividendes,
        "Variation des concours bancaires courants": var_c(["519"]),
    }
    alertes = []
    if prix_cession or db(dn, ["675"]):
        alertes.append("Cessions détectées : la valeur brute sortie est reconstituée (VNC + amortissements cédés). "
                       "À rapprocher du tableau des immobilisations.")
    ouverture, cloture = tresorerie(d1), tresorerie(dn)
    variation = sum(v.values())
    ecart = round(variation - (cloture - ouverture), 2)
    if abs(ecart) >= 1:
        alertes.append(f"Contrôle : la variation calculée diffère de {nb_fr(ecart, 0)} € de la variation réelle "
                       "de trésorerie (comptes non classés ou balances incohérentes).".replace(",", " "))
    v["_treso_ouverture"] = ouverture
    v["_treso_cloture"] = cloture
    v["_alertes"] = alertes
    return v


# ─── Graphique ───────────────────────────────────────────────────────────────

def _chart_tft(resultats: dict, exercices: list) -> go.Figure:
    labels = ["Activité", "Investissement", "Financement"]
    keys = ["Flux activité (I)", "Flux investissement (II)", "Flux financement (III)"]
    colors = ["#1E8449", "#2980B9", "#8E44AD"]

    fig = go.Figure()
    for label, key, color in zip(labels, keys, colors):
        fig.add_trace(go.Bar(
            name=label, x=exercices,
            y=[resultats[ex][key] for ex in exercices],
            marker_color=color,
            hovertemplate="%{y:+,.0f} €<extra>" + label + "</extra>"
        ))
    fig.add_trace(go.Scatter(
        name="Trésorerie clôture", x=exercices,
        y=[resultats[ex]["Trésorerie clôture"] for ex in exercices],
        mode="lines+markers",
        line=dict(color="#E74C3C", width=2.5),
        marker=dict(size=9),
        hovertemplate="%{y:,.0f} €<extra>Trésorerie clôture</extra>"
    ))
    fig.add_hline(y=0, line_dash="dash", line_color="#888", line_width=1)
    fig.update_layout(
        title="Tableau de Flux de Trésorerie — Méthode indirecte",
        barmode="group", height=420,
        margin=dict(l=20, r=20, t=50, b=20),
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
        legend=dict(orientation="h", y=-0.15),
        yaxis_tickformat=",.0f",
    )
    return fig


# ─── Export Excel ────────────────────────────────────────────────────────────

def _export_excel_tft(data: dict, resultats: dict, exercices: list, entreprise: str) -> bytes:
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        rows = []
        for section, info in TFT_STRUCTURE.items():
            rows.append({"Libellé": section, **{ex: "" for ex in exercices}, "_type": "section"})
            for lib, signe in info["lignes"]:
                row = {"Libellé": f"  {lib} ({signe})", "_type": "ligne"}
                for ex in exercices:
                    row[ex] = data.get(lib, {}).get(ex, 0)
                rows.append(row)

        rows.append({"Libellé": "FLUX ACTIVITÉ (I)", **{ex: resultats[ex]["Flux activité (I)"] for ex in exercices}, "_type": "total"})
        rows.append({"Libellé": "FLUX INVESTISSEMENT (II)", **{ex: resultats[ex]["Flux investissement (II)"] for ex in exercices}, "_type": "total"})
        rows.append({"Libellé": "FLUX FINANCEMENT (III)", **{ex: resultats[ex]["Flux financement (III)"] for ex in exercices}, "_type": "total"})
        rows.append({"Libellé": "VARIATION NETTE (I+II+III)", **{ex: resultats[ex]["Variation nette (I+II+III)"] for ex in exercices}, "_type": "total"})
        rows.append({"Libellé": "Trésorerie ouverture", **{ex: resultats[ex]["Trésorerie ouverture"] for ex in exercices}, "_type": "ligne"})
        rows.append({"Libellé": "TRÉSORERIE CLÔTURE", **{ex: resultats[ex]["Trésorerie clôture"] for ex in exercices}, "_type": "total"})

        df_export = pd.DataFrame(rows).drop(columns=["_type"])
        df_export.to_excel(writer, sheet_name="TFT", index=False)

        ws = writer.sheets["TFT"]
        types = [r["_type"] for r in rows]
        section_fill = PatternFill("solid", fgColor="2C3E50")
        total_fill   = PatternFill("solid", fgColor="1E8449")

        for i, t in enumerate(types, start=2):
            if t == "section":
                for cell in ws[i]:
                    cell.fill = section_fill
                    cell.font = Font(color="FFFFFF", bold=True)
            elif t == "total":
                for cell in ws[i]:
                    cell.fill = total_fill
                    cell.font = Font(color="FFFFFF", bold=True)

        for col in ws.columns:
            ws.column_dimensions[get_column_letter(col[0].column)].width = 35

    return buf.getvalue()


# ─── Analyse IA ──────────────────────────────────────────────────────────────

def _analyser_ia(resultats: dict, exercices: list, entreprise: str) -> str:
    try:
        from utils.ai import appel_mistral, extraire_contenu_mistral
        lignes = []
        for ex in exercices:
            r = resultats[ex]
            lignes.append(
                f"  {ex} : Activité={nb_fr_signe(r['Flux activité (I)'], 0)}€ | "
                f"Investissement={nb_fr_signe(r['Flux investissement (II)'], 0)}€ | "
                f"Financement={nb_fr_signe(r['Flux financement (III)'], 0)}€ | "
                f"Tréso clôture={nb_fr(r['Trésorerie clôture'], 0)}€"
            )
        prompt = f"""Tu es expert-comptable PCG France. Analyse ce TFT (méthode indirecte) pour {entreprise} :

{chr(10).join(lignes)}

Fournis :
1. Diagnostic de la santé de trésorerie
2. Qualité des flux d'activité (autofinancement)
3. Politique d'investissement (croissance ou désinvestissement)
4. Structure de financement (endettement, fonds propres)
5. Risques de liquidité identifiés
6. Recommandations concrètes

Sois concis et professionnel."""
        result = appel_mistral(prompt, temperature=0.3,
                               noms=[entreprise] if entreprise and entreprise.strip().lower() != "entreprise" else [])
        return extraire_contenu_mistral(result) or "Analyse indisponible."
    except Exception as e:
        return f"Analyse IA indisponible : {e}"


# ─── Page principale ─────────────────────────────────────────────────────────

def exercices_renseignes(resultats: dict, exercices: list) -> list:
    """Exercices comportant au moins un montant (le dernier est toujours gardé)."""
    utiles = [ex for ex in exercices if any(abs(v) > 0.005 for v in resultats[ex].values())]
    return utiles or exercices[-1:]


def rapport_tft(data: dict, resultats: dict, exercices: list, entreprise: str) -> str:
    """Rapport du TFT (Markdown) : flux par section, synthèse, lecture.
    Seuls les exercices renseignés sont repris (un exercice sans aucun montant n'est pas un exercice à flux nuls)."""
    from datetime import datetime
    exercices = exercices_renseignes(resultats, exercices)
    r = [f"# TABLEAU DES FLUX DE TRÉSORERIE – {entreprise}",
         f"## Exercice{'s' if len(exercices) > 1 else ''} {', '.join(exercices)}",
         f"*Méthode indirecte, modèle OEC · édité le {datetime.now().strftime('%d/%m/%Y')} · montants en euros*",
         "", "---", ""]
    noms_flux = {K_OP: "Flux activité (I)", K_INV: "Flux investissement (II)", K_FIN: "Flux financement (III)"}
    for section, info in TFT_STRUCTURE.items():
        num, _, titre = section.partition(". ")
        r += [f"## {num}. {titre[:1]}{titre[1:].lower()}", "", "| Poste | " + " | ".join(exercices) + " |",
              "|---|" + "---:|" * len(exercices)]
        for lib, _ in info["lignes"]:
            vals = [float(data.get(lib, {}).get(ex, 0) or 0) for ex in exercices]
            if any(vals):
                r.append(f"| {lib} | " + " | ".join(nb_fr(v) for v in vals) + " |")
        tot = [resultats[ex][noms_flux[section]] for ex in exercices]
        r += [f"| **Total** | " + " | ".join(f"**{nb_fr(v)}**" for v in tot) + " |", ""]
    r += ["## Synthèse", "", "| | " + " | ".join(exercices) + " |", "|---|" + "---:|" * len(exercices)]
    for cle in ("Flux activité (I)", "Flux investissement (II)", "Flux financement (III)",
                "Variation nette (I+II+III)", "Trésorerie ouverture", "Trésorerie clôture"):
        g = "**" if cle in ("Variation nette (I+II+III)", "Trésorerie clôture") else ""
        r.append(f"| {g}{cle}{g} | " + " | ".join(f"{g}{nb_fr(resultats[ex][cle])}{g}" for ex in exercices) + " |")
    d = resultats[exercices[-1]]
    r += ["", "## Lecture", ""]
    r.append(f"- **Activité** : {'génère' if d['Flux activité (I)'] >= 0 else 'consomme'} "
             f"{nb_fr(abs(d['Flux activité (I)']))} € de trésorerie en {exercices[-1]}.")
    if d['Flux activité (I)'] + d['Flux investissement (II)'] < 0:
        r.append("- **Les investissements ne sont pas couverts par l'activité** : ils sont financés par la trésorerie "
                 "existante ou par le financement externe.")
    if d['Trésorerie clôture'] < 0:
        r.append("- **Trésorerie de clôture négative** : besoin de financement à court terme.")
    r += ["", "---", "*SMD Global Consulting LLC - Superviseur IA Comptable*"]
    return "\n".join(r)


def visuels_tft(resultats: dict, exercices: list):
    """Indicateurs et graphique pour l'export Word (mêmes chiffres qu'à l'écran)."""
    from utils.word_visuels import barres_et_courbe, COULEUR_N, COULEUR_N1, COULEUR_ORANGE
    exercices = exercices_renseignes(resultats, exercices)
    d = resultats[exercices[-1]]
    signe = lambda x: ("+" if x > 0 else "") + nb_fr(x) + " €"
    ind = [{"libelle": f"Flux d'activité {exercices[-1]}", "valeur": signe(d['Flux activité (I)']),
            "ton": "bon" if d['Flux activité (I)'] >= 0 else "mauvais", "detail": ""},
           {"libelle": "Flux d'investissement", "valeur": signe(d['Flux investissement (II)'])},
           {"libelle": "Flux de financement", "valeur": signe(d['Flux financement (III)'])},
           {"libelle": "Variation de trésorerie", "valeur": signe(d['Variation nette (I+II+III)'])},
           {"libelle": "Trésorerie de clôture", "valeur": nb_fr(d['Trésorerie clôture']) + " €",
            "ton": "bon" if d['Trésorerie clôture'] >= 0 else "mauvais",
            "detail": "positive" if d['Trésorerie clôture'] >= 0 else "négative"}]
    g = barres_et_courbe(exercices,
                         [("Activité", [resultats[e]['Flux activité (I)'] for e in exercices], COULEUR_N),
                          ("Investissement", [resultats[e]['Flux investissement (II)'] for e in exercices], COULEUR_N1),
                          ("Financement", [resultats[e]['Flux financement (III)'] for e in exercices], COULEUR_ORANGE)],
                         ("Trésorerie de clôture", [resultats[e]['Trésorerie clôture'] for e in exercices]),
                         "Flux de trésorerie et trésorerie de clôture")
    return ind, [g]


def page_tft():
    st.title("💹 Tableau de Flux de Trésorerie")
    st.markdown("*Méthode indirecte — modèle OEC — PCG France*")
    st.divider()

    col1, col2, col3 = st.columns([2, 1, 1])
    with col1:
        entreprise = st.text_input("Entreprise", value="Mon Entreprise")
    with col2:
        annee_ref = st.number_input("Exercice de référence", value=None, placeholder="ex. 2025",
                                     min_value=2000, max_value=2050, step=1)
    with col3:
        nb_ex = st.slider("Exercices comparatifs", 1, 3, 2)

    if annee_ref is None:
        st.info("Saisissez l'exercice de référence pour construire le tableau de flux de trésorerie.")
        return
    annee_ref = int(annee_ref)
    exercices = [str(annee_ref - i) for i in range(nb_ex - 1, -1, -1)]
    st.caption(f"Exercices : {' | '.join(exercices)}")
    st.divider()

    # ── Initialisation données
    if "tft_data" not in st.session_state:
        st.session_state.tft_data = {}
    data = st.session_state.tft_data

    # ── Import de deux balances (N et N-1)
    with st.expander("📂 Importer les balances N et N-1 pour pré-remplir", expanded=True):
        st.caption("Balances de clôture avant affectation du résultat (Excel, CSV ou TXT). "
                   "Les variations sont calculées entre les deux exercices.")
        col_n, col_n1, col_ex = st.columns([2, 2, 1])
        with col_n:
            f_n = st.file_uploader("Balance de l'exercice (N)", type=TYPES_BALANCE, key="balance_tft_n")
        with col_n1:
            f_n1 = st.file_uploader("Balance de l'exercice précédent (N-1)", type=TYPES_BALANCE, key="balance_tft_n1")
        with col_ex:
            ex_import = st.selectbox("Exercice N", exercices, index=len(exercices) - 1, key="ex_import_tft")
        if f_n and f_n1 and st.button("Calculer le TFT", key="btn_extract_tft", type="primary"):
            from utils.intelligent_parser import charger_balance_ou_fec
            with st.spinner("Extraction..."):
                df_n, _, _ = charger_balance_ou_fec(f_n)
                df_n1, _, _ = charger_balance_ou_fec(f_n1)
                vals = _extraire_tft(df_n, df_n1)
            for alerte in vals.pop("_alertes"):
                st.warning(alerte)
            data.setdefault("Trésorerie à l'ouverture", {})[ex_import] = round(vals.pop("_treso_ouverture"), 2)
            vals.pop("_treso_cloture")
            for lib, val in vals.items():
                data.setdefault(lib, {})[ex_import] = round(val, 2)
            st.session_state.tft_data = data
            for k in list(st.session_state.keys()):
                if str(k).startswith("tft_"):
                    del st.session_state[k]   # force le rafraîchissement des tableaux
            st.success(f"TFT pré-rempli pour {ex_import}")
        elif f_n and not f_n1:
            st.info("Ajoutez aussi la balance N-1 : un TFT se calcule sur des variations.")

    st.divider()

    # ── Saisie par section
    for section, info in TFT_STRUCTURE.items():
        st.markdown(f"**{section}**")
        rows = []
        for lib, signe in info["lignes"]:
            row = {"Poste": f"{lib}  ({signe})"}
            for ex in exercices:
                row[ex] = data.get(lib, {}).get(ex, 0.0)
            rows.append((lib, row))

        df_section = pd.DataFrame([r for _, r in rows])
        edited = st.data_editor(
            df_section, width="stretch", hide_index=True,
            column_config={ex: st.column_config.NumberColumn(f"{ex} (€)", format="localized")
                           for ex in exercices},
            key=f"tft_{section[:15]}"
        )
        for i, (lib, _) in enumerate(rows):
            if lib not in data:
                data[lib] = {}
            for ex in exercices:
                data[lib][ex] = edited.iloc[i][ex]

    # Trésorerie ouverture
    st.markdown("**Trésorerie**")
    treso_rows = [{"Poste": "Trésorerie à l'ouverture"}]
    for ex in exercices:
        treso_rows[0][ex] = data.get("Trésorerie à l'ouverture", {}).get(ex, 0.0)
    df_treso = pd.DataFrame(treso_rows)
    edited_t = st.data_editor(df_treso, width="stretch", hide_index=True,
                               column_config={ex: st.column_config.NumberColumn(f"{ex} (€)", format="localized")
                                              for ex in exercices},
                               key="tft_treso")
    if "Trésorerie à l'ouverture" not in data:
        data["Trésorerie à l'ouverture"] = {}
    for ex in exercices:
        data["Trésorerie à l'ouverture"][ex] = edited_t.iloc[0][ex]
    st.session_state.tft_data = data

    st.divider()

    # ── Synthèse
    resultats = _calculer_tft(data, exercices)
    st.subheader("📊 Synthèse")
    cols = st.columns(len(exercices))
    for i, ex in enumerate(exercices):
        r = resultats[ex]
        with cols[i]:
            sg = lambda x: ("+" if x > 0 else "") + eur_fr(x)
            st.metric(f"Activité {ex}", sg(r['Flux activité (I)']))
            st.metric(f"Investissement {ex}", sg(r['Flux investissement (II)']))
            st.metric(f"Financement {ex}", sg(r['Flux financement (III)']))
            color = "normal" if r["Trésorerie clôture"] >= 0 else "inverse"
            st.metric(f"Tréso clôture {ex}", eur_fr(r['Trésorerie clôture']), delta_color=color)

    st.plotly_chart(_chart_tft(resultats, exercices), width="stretch")
    st.divider()

    col_ia, col_w, col_xl = st.columns(3)
    with col_w:
        from utils.page_helpers import generer_bouton_word
        ind_w, graph_w = visuels_tft(resultats, exercices)
        generer_bouton_word(f"TFT_{entreprise}_{exercices[-1]}", rapport_tft(data, resultats, exercices, entreprise),
                            indicateurs=ind_w, graphiques=graph_w)
    with col_ia:
        if st.button("🤖 Analyse IA", type="primary", width="stretch"):
            with st.spinner("Analyse en cours..."):
                analyse = _analyser_ia(resultats, exercices, entreprise)
            st.markdown("### 🤖 Analyse IA")
            from utils.page_helpers import afficher_contenu_ia
            afficher_contenu_ia(analyse, "tft")
    with col_xl:
        excel_bytes = _export_excel_tft(data, resultats, exercices, entreprise)
        st.download_button(
            "📥 Exporter Excel",
            data=excel_bytes,
            file_name=sanitize_filename(f"TFT_{entreprise}_{exercices[-1]}") + ".xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            width="stretch",
        )
